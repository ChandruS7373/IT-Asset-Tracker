import os
import json
import sqlite3
import random
import smtplib
import configparser
from datetime import datetime, timedelta
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, send_file
import openpyxl
import io

app = Flask(__name__)
app.secret_key = 'ittracker_secret_2024'

DB_PATH = os.path.join(os.path.dirname(__file__), 'data', 'ittracker.db')
EXCEL_PATH = r'C:\Users\YashwanthHP\Downloads\IT Asset working (1).xlsx'
CONFIG_PATH = os.path.join(os.path.dirname(__file__), 'email_config.ini')


def get_smtp_config():
    cfg = configparser.ConfigParser()
    cfg.read(CONFIG_PATH)
    return cfg['SMTP'] if 'SMTP' in cfg else {}


def get_notification_prefs():
    cfg = configparser.ConfigParser()
    cfg.read(CONFIG_PATH)
    sec = cfg['NOTIFICATIONS'] if 'NOTIFICATIONS' in cfg else {}
    def b(key):
        return sec.get(key, 'true').lower() != 'false'
    return {
        'new_request':          b('notify_new_request'),
        'employee_action':      b('notify_employee_action'),
        'technician_assigned':  b('notify_technician_assigned'),
        'status_employee':      b('notify_status_employee'),
        'status_admin':         b('notify_status_admin'),
        'chat':                 b('notify_chat'),
    }


def send_otp_email(to_email, otp_code, employee_name):
    cfg = get_smtp_config()
    host = cfg.get('SMTP_HOST', '')
    port = int(cfg.get('SMTP_PORT', 587))
    user = cfg.get('SMTP_USER', '')
    pwd  = cfg.get('SMTP_PASSWORD', '')
    from_name = cfg.get('FROM_NAME', 'Qualesce IT Tracker')

    if not user or user == 'your-email@gmail.com':
        return False, 'SMTP not configured. Edit email_config.ini with your credentials.'

    msg = MIMEMultipart('alternative')
    msg['Subject'] = 'IT Asset Tracker - Password Reset Code'
    msg['From']    = f'{from_name} <{user}>'
    msg['To']      = to_email

    html = f"""
    <div style="font-family:Segoe UI,sans-serif;max-width:480px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
      <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:28px 30px;text-align:center">
        <h2 style="color:#fff;margin:0;font-size:1.3rem">Password Reset Request</h2>
      </div>
      <div style="padding:30px">
        <p style="color:#374151">Hi <strong>{employee_name}</strong>,</p>
        <p style="color:#374151">Use the code below to reset your password. It expires in <strong>10 minutes</strong>.</p>
        <div style="text-align:center;margin:28px 0">
          <span style="background:#f0f4ff;border:2px dashed #6366f1;border-radius:10px;padding:16px 36px;font-size:2.2rem;font-weight:800;letter-spacing:10px;color:#1e1b4b">{otp_code}</span>
        </div>
        <p style="color:#6b7280;font-size:0.85rem">If you did not request this, ignore this email. Your password will not change.</p>
      </div>
      <div style="background:#f8fafc;padding:14px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">
        Qualesce IT Asset Tracker &nbsp;|&nbsp; Do not reply to this email
      </div>
    </div>
    """
    msg.attach(MIMEText(html, 'html'))

    try:
        with smtplib.SMTP(host, port, timeout=10) as s:
            s.ehlo()
            s.starttls()
            s.login(user, pwd)
            s.sendmail(user, to_email, msg.as_string())
        return True, 'OTP sent successfully'
    except Exception as e:
        return False, str(e)

def get_user_notif_prefs(user_id):
    conn = get_db()
    row = conn.execute("SELECT notification_prefs FROM users WHERE id=?", (user_id,)).fetchone()
    conn.close()
    try:
        return json.loads(row['notification_prefs']) if row and row['notification_prefs'] else {}
    except Exception:
        return {}


def _should_notify_employee(employee_id, emp_email, global_key, pref_key):
    if not emp_email:
        return False
    if not get_notification_prefs().get(global_key, True):
        return False
    return get_user_notif_prefs(employee_id).get(pref_key, True)


def _should_notify_tech(tech_id, tech_email, global_key, pref_key):
    if not tech_email:
        return False
    if not get_notification_prefs().get(global_key, True):
        return False
    return get_user_notif_prefs(tech_id).get(pref_key, True)


def send_notification_email(to_emails, subject, html_body):
    cfg = get_smtp_config()
    host = cfg.get('smtp_host', '')
    port = int(cfg.get('smtp_port', 587))
    user = cfg.get('smtp_user', '')
    pwd  = cfg.get('smtp_password', '')
    from_name = cfg.get('from_name', 'Qualesce IT Tracker')
    if not user or user == 'your-email@gmail.com':
        return False, 'SMTP not configured'
    if isinstance(to_emails, str):
        to_emails = [to_emails]
    to_emails = [e for e in to_emails if e]
    if not to_emails:
        return False, 'No recipients'
    msg = MIMEMultipart('alternative')
    msg['Subject'] = subject
    msg['From']    = f'{from_name} <{user}>'
    msg['To']      = ', '.join(to_emails)
    msg.attach(MIMEText(html_body, 'html'))
    try:
        with smtplib.SMTP(host, port, timeout=10) as s:
            s.ehlo()
            s.starttls()
            s.login(user, pwd)
            s.sendmail(user, to_emails, msg.as_string())
        return True, 'Email sent'
    except Exception as e:
        return False, str(e)


def build_chat_email_html(sender_name, sender_role, message, req_id, service_type):
    role_labels = {'admin': 'Admin', 'technician': 'Technician', 'user': 'Employee'}
    return f"""
    <div style="font-family:Segoe UI,sans-serif;max-width:560px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
      <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:24px 30px;text-align:center">
        <h2 style="color:#fff;margin:0;font-size:1.2rem">New Message — Request #{req_id}</h2>
      </div>
      <div style="padding:28px 30px">
        <p style="color:#374151">New message from <strong>{sender_name}</strong> ({role_labels.get(sender_role, sender_role)}) on <strong>Request #{req_id}</strong> ({service_type}):</p>
        <div style="background:#f1f5f9;border-left:4px solid #533483;border-radius:4px;padding:14px 18px;margin:18px 0;color:#1e293b">{message}</div>
        <p style="color:#6b7280;font-size:0.9rem">Please log in to IT Asset Tracker to reply.</p>
      </div>
      <div style="background:#f8fafc;padding:12px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">Qualesce IT Asset Tracker</div>
    </div>"""


def build_status_email_html(employee_name, new_status, req_id, service_type, asset_no):
    colors = {'In Progress': '#f59e0b', 'Hold': '#6b7280', 'Completed': '#22c55e',
              'Accepted': '#22c55e', 'Declined': '#ef4444', 'Pending': '#3b82f6'}
    color = colors.get(new_status, '#374151')
    return f"""
    <div style="font-family:Segoe UI,sans-serif;max-width:560px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
      <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:24px 30px;text-align:center">
        <h2 style="color:#fff;margin:0;font-size:1.2rem">Request #{req_id} — Status Updated</h2>
      </div>
      <div style="padding:28px 30px">
        <p style="color:#374151">Hi <strong>{employee_name}</strong>,</p>
        <p style="color:#374151">Your service request status is now <strong style="color:{color}">{new_status}</strong>.</p>
        <table style="width:100%;border-collapse:collapse;margin:16px 0">
          <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;width:40%;font-weight:600">Request #</td><td style="padding:8px 12px;color:#1e293b">#{req_id}</td></tr>
          <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Asset</td><td style="padding:8px 12px;color:#1e293b">{asset_no}</td></tr>
          <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;font-weight:600">Issue Type</td><td style="padding:8px 12px;color:#1e293b">{service_type}</td></tr>
          <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">New Status</td><td style="padding:8px 12px;font-weight:700;color:{color}">{new_status}</td></tr>
        </table>
      </div>
      <div style="background:#f8fafc;padding:12px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">Qualesce IT Asset Tracker</div>
    </div>"""


# ---------------------------------------------------------------------------
# Database setup
# ---------------------------------------------------------------------------

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    c = conn.cursor()
    c.executescript('''
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password TEXT NOT NULL,
            role TEXT NOT NULL DEFAULT 'user',
            employee_name TEXT,
            email TEXT,
            department TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS assets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_no TEXT UNIQUE,
            asset_status TEXT,
            login_id TEXT,
            serial_no TEXT,
            transfer_history TEXT,
            model TEXT,
            years TEXT,
            resolution TEXT,
            sn TEXT,
            model2 TEXT,
            warranty_start TEXT,
            warranty_end TEXT,
            warranty_type TEXT,
            warranty_status TEXT,
            lan_mac TEXT,
            lan_ip TEXT,
            wireless_mac TEXT,
            wan_ip TEXT,
            admin_usb TEXT,
            bios_password TEXT,
            admin_password TEXT,
            spiceworks TEXT,
            windows_update TEXT,
            unwanted_apps TEXT,
            processor TEXT,
            ram TEXT,
            hdd TEXT,
            office365 TEXT,
            sharepoint TEXT,
            onedrive TEXT,
            worksoft TEXT,
            ia TEXT,
            sql_version TEXT,
            system_cleanup TEXT,
            created_at TEXT DEFAULT (datetime('now')),
            updated_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS service_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_no TEXT,
            service_date TEXT,
            service_type TEXT,
            description TEXT,
            technician TEXT,
            status TEXT DEFAULT 'Completed',
            cost TEXT,
            remarks TEXT,
            logged_by TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS machine_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_no TEXT,
            employee_name TEXT,
            assigned_date TEXT,
            returned_date TEXT,
            remarks TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS password_resets (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT NOT NULL,
            otp TEXT NOT NULL,
            expires_at TEXT NOT NULL,
            used INTEGER DEFAULT 0,
            created_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS stock_dashboard (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            status TEXT NOT NULL,
            model TEXT NOT NULL,
            count INTEGER DEFAULT 0,
            UNIQUE(status, model)
        );

        CREATE TABLE IF NOT EXISTS service_requests (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            asset_no TEXT,
            employee_id INTEGER,
            employee_name TEXT,
            service_type TEXT,
            description TEXT,
            remarks TEXT,
            status TEXT DEFAULT 'Pending',
            assigned_to_id INTEGER,
            assigned_to_name TEXT,
            created_at TEXT DEFAULT (datetime('now')),
            updated_at TEXT DEFAULT (datetime('now'))
        );

        CREATE TABLE IF NOT EXISTS chat_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            request_id INTEGER,
            sender_id INTEGER,
            sender_name TEXT,
            sender_role TEXT,
            message TEXT,
            created_at TEXT DEFAULT (datetime('now'))
        );
    ''')

    # Default admin
    c.execute("SELECT id FROM users WHERE username='admin'")
    if not c.fetchone():
        c.execute("INSERT INTO users (username, password, role, employee_name) VALUES (?, ?, ?, ?)",
                  ('admin', 'admin123', 'admin', 'Administrator'))

    # Seed initial stock dashboard data
    c.execute("SELECT COUNT(*) FROM stock_dashboard")
    if c.fetchone()[0] == 0:
        seed = [
            ('Assigned',  'Dell LATITUDE 7480',     3),
            ('DEAD',      'DELL Inspiron N5010',     1),
            ('DEAD',      'Dell LATITUDE 7480',      2),
            ('DEAD',      'DELL Vostro 3558',        8),
            ('DEAD',      'Lenovo B40-80',           5),
            ('DEAD',      'LENOVO V310',             2),
            ('IT Stock',  'DELL LATITUDE 3400',      3),
            ('IT Stock',  'DELL LATITUDE 3410',      2),
            ('IT Stock',  'Dell LATITUDE 7480',      6),
            ('IT Stock',  'DELL Vostro 3558',        1),
            ('IT Stock',  'Lenovo B40-80',           4),
            ('IT Stock',  'LENOVO V130',             2),
            ('IT Stock',  'LENOVO V310',            12),
            ('Service',   'Dell LATITUDE 7480',      3),
            ('To Check',  'DELL LATITUDE 3400',      4),
            ('To Check',  'DELL LATITUDE 3410',      8),
            ('To Check',  'Dell Latitude 5310',      1),
            ('To Check',  'Dell LATITUDE 7480',     75),
            ('To Check',  'DELL LATITUDE E7450',     1),
            ('To Check',  'DELL Vostro 3558',        1),
            ('To Check',  'LENOVO THINKPAD T450',   19),
            ('To Check',  'Lenovo B40-80',           1),
            ('To Check',  'LENOVO V130',             8),
            ('To Check',  'LENOVO V310',            18),
        ]
        c.executemany("INSERT OR IGNORE INTO stock_dashboard (status, model, count) VALUES (?,?,?)", seed)

    # Add notification_prefs column if not present (migration)
    try:
        c.execute("ALTER TABLE users ADD COLUMN notification_prefs TEXT DEFAULT '{}'")
    except Exception:
        pass

    conn.commit()
    conn.close()


def import_excel():
    if not os.path.exists(EXCEL_PATH):
        return
    conn = get_db()
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM assets")
    if c.fetchone()[0] > 0:
        conn.close()
        return  # already imported

    try:
        wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True)
        ws = wb['Hardware']
        imported = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            def s(v): return str(v).strip() if v is not None else ''
            def d(v):
                if v is None: return ''
                try:
                    return str(v.date()) if hasattr(v, 'date') else str(v)
                except:
                    return str(v)

            asset_no = s(row[1])
            if not asset_no:
                continue
            # Skip phantom/garbage rows: asset_no exists but no meaningful data
            if not row[0] and not row[2] and not row[5]:
                continue

            c.execute('''INSERT OR IGNORE INTO assets
                (asset_status, asset_no, login_id, serial_no, transfer_history,
                 model, years, resolution, sn, model2,
                 warranty_start, warranty_end, warranty_type, warranty_status,
                 lan_mac, lan_ip, wireless_mac, wan_ip,
                 admin_usb, bios_password, admin_password,
                 spiceworks, windows_update, unwanted_apps,
                 processor, ram, hdd, office365, sharepoint, onedrive,
                 worksoft, ia, sql_version, system_cleanup)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (s(row[0]), asset_no, s(row[2]), s(row[3]), s(row[4]),
                 s(row[5]), s(row[6]), s(row[7]), s(row[8]), s(row[10]),
                 d(row[11]), d(row[12]), s(row[13]), s(row[14]),
                 s(row[15]), s(row[16]), s(row[17]), s(row[18]),
                 s(row[19]), s(row[20]), s(row[21]),
                 s(row[22]), s(row[23]), s(row[24]),
                 s(row[25]), s(row[26]), s(row[27]), s(row[28]),
                 s(row[29]), s(row[30]), s(row[31]), s(row[32]),
                 s(row[33]), s(row[37]) if len(row) > 37 else ''))

            # Parse machine history from transfer_history
            transfer = s(row[4])
            if transfer:
                names = [n.strip() for n in transfer.replace(' - ', '-').split('-') if n.strip()]
                for name in names:
                    c.execute('''INSERT INTO machine_history (asset_no, employee_name, remarks)
                                 VALUES (?, ?, ?)''', (asset_no, name, 'Imported from Excel'))

            # Create user account for assigned employee
            login_id = s(row[2])
            if login_id and login_id not in ('IT Stock', 'DEAD', ''):
                c.execute("SELECT id FROM users WHERE employee_name=?", (login_id,))
                if not c.fetchone():
                    username = login_id.lower().replace(' ', '.').replace('/', '')[:20]
                    c.execute("SELECT id FROM users WHERE username=?", (username,))
                    if not c.fetchone():
                        c.execute("INSERT INTO users (username, password, role, employee_name) VALUES (?,?,?,?)",
                                  (username, 'pass123', 'user', login_id))
            imported += 1

        wb.close()
        conn.commit()
        print(f"Imported {imported} assets from Excel")
    except Exception as e:
        print(f"Excel import error: {e}")
    finally:
        conn.close()


# ---------------------------------------------------------------------------
# Auth helpers
# ---------------------------------------------------------------------------

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated


def admin_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if session.get('role') != 'admin':
            flash('Admin access required.', 'danger')
            return redirect(url_for('user_dashboard'))
        return f(*args, **kwargs)
    return decorated


def technician_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if session.get('role') != 'technician':
            flash('Technician access required.', 'danger')
            if session.get('role') == 'admin':
                return redirect(url_for('admin_dashboard'))
            return redirect(url_for('user_dashboard'))
        return f(*args, **kwargs)
    return decorated


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------

@app.route('/')
def index():
    if 'user_id' in session:
        role = session.get('role')
        if role == 'admin':
            return redirect(url_for('admin_dashboard'))
        if role == 'technician':
            return redirect(url_for('technician_dashboard'))
        return redirect(url_for('user_dashboard'))
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        email    = request.form['email'].strip().lower()
        password = request.form['password'].strip()
        conn = get_db()
        # Match by email or username (fallback for admin)
        user = conn.execute(
            "SELECT * FROM users WHERE (LOWER(email)=? OR LOWER(username)=?) AND password=?",
            (email, email, password)).fetchone()
        conn.close()
        if user:
            session['user_id']       = user['id']
            session['username']      = user['username']
            session['role']          = user['role']
            session['employee_name'] = user['employee_name'] or user['username']
            session['user_email']    = user['email'] or ''
            if user['role'] == 'admin':
                return redirect(url_for('admin_dashboard'))
            if user['role'] == 'technician':
                return redirect(url_for('technician_dashboard'))
            return redirect(url_for('user_dashboard'))
        flash('Invalid email or password.', 'danger')
    return render_template('login.html')


@app.route('/forgot-password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        email = request.form['email'].strip().lower()
        conn  = get_db()
        user  = conn.execute("SELECT * FROM users WHERE LOWER(email)=?", (email,)).fetchone()
        if not user:
            conn.close()
            flash('No account found with that email address.', 'danger')
            return render_template('forgot_password.html')

        otp      = str(random.randint(100000, 999999))
        expires  = (datetime.now() + timedelta(minutes=10)).isoformat()
        # Invalidate previous OTPs for this email
        conn.execute("UPDATE password_resets SET used=1 WHERE email=?", (email,))
        conn.execute("INSERT INTO password_resets (email, otp, expires_at) VALUES (?,?,?)",
                     (email, otp, expires))
        conn.commit()
        conn.close()

        ok, msg = send_otp_email(email, otp, user['employee_name'] or user['username'])
        if ok:
            session['reset_email'] = email
            flash(f'A 6-digit code has been sent to {email}', 'success')
            return redirect(url_for('verify_otp'))
        else:
            # Dev fallback: show OTP on screen if SMTP not configured
            flash(f'Email not configured — use this code to reset: {otp}', 'warning')
            session['reset_email'] = email
            return redirect(url_for('verify_otp'))

    return render_template('forgot_password.html')


@app.route('/verify-otp', methods=['GET', 'POST'])
def verify_otp():
    email = session.get('reset_email')
    if not email:
        return redirect(url_for('forgot_password'))

    if request.method == 'POST':
        entered = request.form['otp'].strip()
        conn    = get_db()
        record  = conn.execute(
            "SELECT * FROM password_resets WHERE email=? AND used=0 ORDER BY id DESC LIMIT 1",
            (email,)).fetchone()

        if not record:
            conn.close()
            flash('No active reset code found. Please request again.', 'danger')
            return redirect(url_for('forgot_password'))

        if datetime.fromisoformat(record['expires_at']) < datetime.now():
            conn.execute("UPDATE password_resets SET used=1 WHERE id=?", (record['id'],))
            conn.commit()
            conn.close()
            flash('Code has expired. Please request a new one.', 'danger')
            return redirect(url_for('forgot_password'))

        if record['otp'] != entered:
            conn.close()
            flash('Incorrect code. Please try again.', 'danger')
            return render_template('verify_otp.html', email=email)

        # Mark used and allow reset
        conn.execute("UPDATE password_resets SET used=1 WHERE id=?", (record['id'],))
        conn.commit()
        conn.close()
        session['reset_verified'] = True
        return redirect(url_for('reset_password'))

    return render_template('verify_otp.html', email=email)


@app.route('/reset-password', methods=['GET', 'POST'])
def reset_password():
    email = session.get('reset_email')
    if not email or not session.get('reset_verified'):
        return redirect(url_for('forgot_password'))

    if request.method == 'POST':
        new_pass  = request.form['password'].strip()
        confirm   = request.form['confirm'].strip()
        if new_pass != confirm:
            flash('Passwords do not match.', 'danger')
            return render_template('reset_password.html')
        if len(new_pass) < 6:
            flash('Password must be at least 6 characters.', 'danger')
            return render_template('reset_password.html')

        conn = get_db()
        conn.execute("UPDATE users SET password=? WHERE LOWER(email)=?", (new_pass, email))
        conn.commit()
        conn.close()

        session.pop('reset_email', None)
        session.pop('reset_verified', None)
        flash('Password updated successfully! Please log in.', 'success')
        return redirect(url_for('login'))

    return render_template('reset_password.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))


# ---------------------------------------------------------------------------
# Admin routes
# ---------------------------------------------------------------------------

@app.route('/admin')
@admin_required
def admin_dashboard():
    conn = get_db()
    # Count per status dynamically
    status_counts = {r[0]: r[1] for r in conn.execute(
        "SELECT asset_status, COUNT(*) FROM assets GROUP BY asset_status").fetchall()}
    stats = {
        'total':    sum(status_counts.values()),
        'assigned': status_counts.get('Assigned', 0),
        'stock':    status_counts.get('IT Stock', 0) + status_counts.get('ITStock', 0),
        'dead':     status_counts.get('DEAD', 0),
        'to_check': status_counts.get('To Check', 0),
        'yashwanth': status_counts.get('Yashwanth', 0),
        'users':    conn.execute("SELECT COUNT(*) FROM users WHERE role='user'").fetchone()[0],
        'services': conn.execute("SELECT COUNT(*) FROM service_history").fetchone()[0],
    }
    recent = conn.execute(
        "SELECT * FROM assets ORDER BY updated_at DESC LIMIT 10").fetchall()

    # Build stock dashboard pivot
    stock_rows = conn.execute(
        "SELECT status, model, count FROM stock_dashboard ORDER BY model").fetchall()
    conn.close()

    stock_models = []
    stock_statuses_set = []
    stock_data = {}
    for row in stock_rows:
        if row['model'] not in stock_models:
            stock_models.append(row['model'])
        if row['status'] not in stock_statuses_set:
            stock_statuses_set.append(row['status'])
        stock_data.setdefault(row['status'], {})[row['model']] = row['count']

    # Standard statuses first, then any extras from actual data
    standard = ['Assigned', 'IT Stock', 'Service', 'To Check', 'DEAD']
    stock_statuses = [s for s in standard if s in stock_statuses_set] + \
                     [s for s in stock_statuses_set if s not in standard]

    return render_template('admin/dashboard.html', stats=stats, recent=recent,
                           stock_models=stock_models, stock_data=stock_data,
                           stock_statuses=stock_statuses)


@app.route('/admin/assets')
@admin_required
def admin_assets():
    search = request.args.get('q', '')
    status = request.args.get('status', '')
    conn = get_db()
    query = "SELECT * FROM assets WHERE 1=1"
    params = []
    if search:
        query += " AND (login_id LIKE ? OR asset_no LIKE ? OR model LIKE ? OR serial_no LIKE ?)"
        params += [f'%{search}%'] * 4
    if status:
        query += " AND asset_status=?"
        params.append(status)
    query += " ORDER BY asset_no"
    assets = conn.execute(query, params).fetchall()
    conn.close()
    return render_template('admin/assets.html', assets=assets, search=search, status=status)


@app.route('/admin/employee/<login_id>')
@admin_required
def admin_employee_detail(login_id):
    conn = get_db()
    assets = conn.execute(
        "SELECT * FROM assets WHERE login_id=? ORDER BY asset_no", (login_id,)).fetchall()
    service_history = conn.execute(
        """SELECT sh.*, a.model FROM service_history sh
           JOIN assets a ON sh.asset_no = a.asset_no
           WHERE a.login_id=? ORDER BY sh.service_date DESC""", (login_id,)).fetchall()
    machine_hist = conn.execute(
        """SELECT mh.* FROM machine_history mh
           JOIN assets a ON mh.asset_no = a.asset_no
           WHERE a.login_id=? ORDER BY mh.created_at DESC""", (login_id,)).fetchall()
    # All history for this employee's name across all machines
    all_machine_hist = conn.execute(
        "SELECT mh.*, a.model FROM machine_history mh LEFT JOIN assets a ON mh.asset_no=a.asset_no WHERE mh.employee_name LIKE ? ORDER BY mh.created_at DESC",
        (f'%{login_id}%',)).fetchall()
    conn.close()
    return render_template('admin/employee_detail.html',
                           login_id=login_id,
                           assets=assets,
                           service_history=service_history,
                           machine_history=all_machine_hist)


@app.route('/admin/asset/<asset_no>')
@admin_required
def admin_asset_detail(asset_no):
    conn = get_db()
    asset = conn.execute("SELECT * FROM assets WHERE asset_no=?", (asset_no,)).fetchone()
    service_history = conn.execute(
        "SELECT * FROM service_history WHERE asset_no=? ORDER BY service_date DESC", (asset_no,)).fetchall()
    machine_history = conn.execute(
        "SELECT * FROM machine_history WHERE asset_no=? ORDER BY created_at DESC", (asset_no,)).fetchall()
    conn.close()
    if not asset:
        flash('Asset not found.', 'danger')
        return redirect(url_for('admin_assets'))
    return render_template('admin/asset_detail.html',
                           asset=asset,
                           service_history=service_history,
                           machine_history=machine_history)


@app.route('/admin/asset/add', methods=['GET', 'POST'])
@admin_required
def admin_add_asset():
    if request.method == 'POST':
        f = request.form
        conn = get_db()
        try:
            conn.execute('''INSERT INTO assets
                (asset_status, asset_no, login_id, serial_no, transfer_history,
                 model, years, resolution, warranty_start, warranty_end,
                 warranty_type, warranty_status, lan_mac, lan_ip,
                 processor, ram, hdd, office365, windows_update, system_cleanup)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (f.get('asset_status'), f.get('asset_no'), f.get('login_id'),
                 f.get('serial_no'), f.get('transfer_history'),
                 f.get('model'), f.get('years'), f.get('resolution'),
                 f.get('warranty_start'), f.get('warranty_end'),
                 f.get('warranty_type'), f.get('warranty_status'),
                 f.get('lan_mac'), f.get('lan_ip'),
                 f.get('processor'), f.get('ram'), f.get('hdd'),
                 f.get('office365'), f.get('windows_update'), f.get('system_cleanup')))
            conn.commit()
            flash('Asset added successfully!', 'success')
            return redirect(url_for('admin_assets'))
        except Exception as e:
            flash(f'Error: {e}', 'danger')
        finally:
            conn.close()
    return render_template('admin/add_asset.html')


@app.route('/admin/asset/edit/<asset_no>', methods=['GET', 'POST'])
@admin_required
def admin_edit_asset(asset_no):
    conn = get_db()
    asset = conn.execute("SELECT * FROM assets WHERE asset_no=?", (asset_no,)).fetchone()
    if not asset:
        flash('Asset not found.', 'danger')
        conn.close()
        return redirect(url_for('admin_assets'))
    if request.method == 'POST':
        f = request.form
        conn.execute('''UPDATE assets SET
            asset_status=?, login_id=?, serial_no=?, transfer_history=?,
            model=?, years=?, resolution=?, warranty_start=?, warranty_end=?,
            warranty_type=?, warranty_status=?, lan_mac=?, lan_ip=?,
            processor=?, ram=?, hdd=?, office365=?, windows_update=?,
            system_cleanup=?, updated_at=?
            WHERE asset_no=?''',
            (f.get('asset_status'), f.get('login_id'), f.get('serial_no'),
             f.get('transfer_history'), f.get('model'), f.get('years'),
             f.get('resolution'), f.get('warranty_start'), f.get('warranty_end'),
             f.get('warranty_type'), f.get('warranty_status'),
             f.get('lan_mac'), f.get('lan_ip'),
             f.get('processor'), f.get('ram'), f.get('hdd'),
             f.get('office365'), f.get('windows_update'), f.get('system_cleanup'),
             datetime.now().isoformat(), asset_no))
        conn.commit()
        conn.close()
        flash('Asset updated!', 'success')
        return redirect(url_for('admin_asset_detail', asset_no=asset_no))
    conn.close()
    return render_template('admin/edit_asset.html', asset=asset)


@app.route('/admin/service/add', methods=['GET', 'POST'])
@admin_required
def admin_add_service():
    conn = get_db()
    if request.method == 'POST':
        f = request.form
        conn.execute('''INSERT INTO service_history
            (asset_no, service_date, service_type, description, technician, status, cost, remarks, logged_by)
            VALUES (?,?,?,?,?,?,?,?,?)''',
            (f.get('asset_no'), f.get('service_date'), f.get('service_type'),
             f.get('description'), f.get('technician'), f.get('status'),
             f.get('cost'), f.get('remarks'), session['employee_name']))
        conn.commit()
        conn.close()
        flash('Service record added!', 'success')
        return redirect(url_for('admin_services'))
    assets = conn.execute("SELECT asset_no, model, login_id FROM assets ORDER BY asset_no").fetchall()
    conn.close()
    return render_template('admin/add_service.html', assets=assets, now=datetime.now().strftime('%Y-%m-%d'))


@app.route('/admin/services')
@admin_required
def admin_services():
    search = request.args.get('q', '')
    conn = get_db()
    query = """SELECT sh.*, a.model, a.login_id FROM service_history sh
               LEFT JOIN assets a ON sh.asset_no=a.asset_no WHERE 1=1"""
    params = []
    if search:
        query += " AND (sh.asset_no LIKE ? OR sh.description LIKE ? OR a.login_id LIKE ?)"
        params += [f'%{search}%'] * 3
    query += " ORDER BY sh.service_date DESC"
    services = conn.execute(query, params).fetchall()
    conn.close()
    return render_template('admin/services.html', services=services, search=search)


@app.route('/admin/users')
@admin_required
def admin_users():
    conn = get_db()
    users = conn.execute("SELECT * FROM users ORDER BY role, username").fetchall()
    conn.close()
    return render_template('admin/users.html', users=users)


@app.route('/admin/users/add', methods=['GET', 'POST'])
@admin_required
def admin_add_user():
    if request.method == 'POST':
        f = request.form
        conn = get_db()
        try:
            conn.execute("INSERT INTO users (username, password, role, employee_name, email, department) VALUES (?,?,?,?,?,?)",
                         (f.get('username'), f.get('password'), f.get('role'),
                          f.get('employee_name'), f.get('email'), f.get('department')))
            conn.commit()
            flash('User created!', 'success')
            return redirect(url_for('admin_users'))
        except Exception as e:
            flash(f'Error: {e}', 'danger')
        finally:
            conn.close()
    return render_template('admin/add_user.html')


@app.route('/admin/users/edit/<int:user_id>', methods=['GET', 'POST'])
@admin_required
def admin_edit_user(user_id):
    conn = get_db()
    user = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if not user:
        conn.close()
        flash('User not found.', 'danger')
        return redirect(url_for('admin_users'))
    if request.method == 'POST':
        f = request.form
        conn.execute(
            "UPDATE users SET employee_name=?, email=?, department=?, role=?, password=? WHERE id=?",
            (f.get('employee_name'), f.get('email'), f.get('department'),
             f.get('role'), f.get('password'), user_id))
        conn.commit()
        conn.close()
        flash('User updated!', 'success')
        return redirect(url_for('admin_users'))
    conn.close()
    return render_template('admin/edit_user.html', user=user)


@app.route('/admin/users/set-email', methods=['POST'])
@admin_required
def admin_set_email_bulk():
    """Quick inline email update from users table."""
    user_id = request.form.get('user_id')
    email   = request.form.get('email', '').strip()
    conn = get_db()
    conn.execute("UPDATE users SET email=? WHERE id=?", (email, user_id))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/admin/employees')
@admin_required
def admin_employees():
    conn = get_db()
    employees = conn.execute(
        """SELECT login_id, COUNT(*) as asset_count,
           GROUP_CONCAT(model, ', ') as models,
           GROUP_CONCAT(asset_no, ', ') as asset_nos,
           MAX(asset_status) as latest_status
           FROM assets
           WHERE login_id NOT IN ('IT Stock', 'DEAD', '')
           GROUP BY login_id
           ORDER BY login_id""").fetchall()
    conn.close()
    return render_template('admin/employees.html', employees=employees)


@app.route('/admin/machine-history/add', methods=['POST'])
@admin_required
def admin_add_machine_history():
    f = request.form
    conn = get_db()
    conn.execute('''INSERT INTO machine_history
        (asset_no, employee_name, assigned_date, returned_date, remarks)
        VALUES (?,?,?,?,?)''',
        (f.get('asset_no'), f.get('employee_name'),
         f.get('assigned_date'), f.get('returned_date'), f.get('remarks')))
    conn.commit()
    conn.close()
    flash('Machine history added!', 'success')
    return redirect(request.referrer or url_for('admin_assets'))


@app.route('/admin/sync-stock', methods=['POST'])
@admin_required
def admin_sync_stock():
    from collections import defaultdict
    conn = get_db()
    rows = conn.execute('''
        SELECT asset_status, model, COUNT(*) as cnt
        FROM assets
        WHERE asset_status IS NOT NULL AND asset_status != ''
          AND model IS NOT NULL AND model != ''
        GROUP BY asset_status, model
    ''').fetchall()
    agg = defaultdict(int)
    for r in rows:
        status = r['asset_status'].strip()
        if status == 'ITStock':
            status = 'IT Stock'
        model = (r['model'] or '').strip()
        if model:
            agg[(status, model)] += r['cnt']
    conn.execute("DELETE FROM stock_dashboard")
    for (status, model), count in agg.items():
        conn.execute("INSERT INTO stock_dashboard (status, model, count) VALUES (?,?,?)",
                     (status, model, count))
    conn.commit()
    conn.close()
    flash('Stock dashboard synced from actual asset data!', 'success')
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/reimport', methods=['POST'])
@admin_required
def admin_reimport():
    conn = get_db()
    conn.execute("DELETE FROM assets")
    conn.execute("DELETE FROM machine_history")
    conn.commit()
    conn.close()
    import_excel()
    flash('Excel data re-imported successfully!', 'success')
    return redirect(url_for('admin_dashboard'))


# ---------------------------------------------------------------------------
# User routes
# ---------------------------------------------------------------------------

@app.route('/user')
@login_required
def user_dashboard():
    if session.get('role') == 'admin':
        return redirect(url_for('admin_dashboard'))
    if session.get('role') == 'technician':
        return redirect(url_for('technician_dashboard'))
    employee_name = session['employee_name']
    conn = get_db()
    assets = conn.execute(
        "SELECT * FROM assets WHERE login_id LIKE ? ORDER BY asset_no",
        (f'%{employee_name}%',)).fetchall()
    service_history = conn.execute(
        """SELECT sh.*, a.model FROM service_history sh
           JOIN assets a ON sh.asset_no=a.asset_no
           WHERE a.login_id LIKE ? ORDER BY sh.service_date DESC LIMIT 10""",
        (f'%{employee_name}%',)).fetchall()
    machine_hist = conn.execute(
        "SELECT mh.*, a.model FROM machine_history mh LEFT JOIN assets a ON mh.asset_no=a.asset_no WHERE mh.employee_name LIKE ? ORDER BY mh.created_at DESC",
        (f'%{employee_name}%',)).fetchall()
    conn.close()
    return render_template('user/dashboard.html',
                           assets=assets,
                           service_history=service_history,
                           machine_history=machine_hist,
                           employee_name=employee_name)


@app.route('/user/asset/<asset_no>')
@login_required
def user_asset_detail(asset_no):
    conn = get_db()
    asset = conn.execute("SELECT * FROM assets WHERE asset_no=?", (asset_no,)).fetchone()
    service_history = conn.execute(
        "SELECT * FROM service_history WHERE asset_no=? ORDER BY service_date DESC",
        (asset_no,)).fetchall()
    machine_history = conn.execute(
        "SELECT * FROM machine_history WHERE asset_no=? ORDER BY created_at DESC",
        (asset_no,)).fetchall()
    conn.close()
    return render_template('user/asset_detail.html',
                           asset=asset,
                           service_history=service_history,
                           machine_history=machine_history)


@app.route('/user/request-service', methods=['GET', 'POST'])
@login_required
def user_request_service():
    if session.get('role') in ('admin', 'technician'):
        return redirect(url_for('admin_dashboard'))
    employee_name = session['employee_name']
    conn = get_db()
    if request.method == 'POST':
        f = request.form
        cur = conn.execute(
            '''INSERT INTO service_requests (asset_no, employee_id, employee_name, service_type, description, remarks, status)
               VALUES (?,?,?,?,?,?,?)''',
            (f.get('asset_no'), session['user_id'], employee_name,
             f.get('service_type'), f.get('description'), f.get('remarks'), 'Pending'))
        req_id = cur.lastrowid
        conn.commit()
        admins = conn.execute("SELECT email FROM users WHERE role='admin' AND email!=''").fetchall()
        admin_emails = [a['email'] for a in admins if a['email']]
        conn.close()
        if admin_emails and get_notification_prefs()['new_request']:
            html = f"""
            <div style="font-family:Segoe UI,sans-serif;max-width:560px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
              <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:24px 30px;text-align:center">
                <h2 style="color:#fff;margin:0;font-size:1.2rem">New Service Request #{req_id}</h2>
              </div>
              <div style="padding:28px 30px">
                <p style="color:#374151">A new service request requires your attention.</p>
                <table style="width:100%;border-collapse:collapse;margin:16px 0">
                  <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;width:40%;font-weight:600">Employee</td><td style="padding:8px 12px;color:#1e293b">{employee_name}</td></tr>
                  <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Asset No</td><td style="padding:8px 12px;color:#1e293b">{f.get('asset_no','-')}</td></tr>
                  <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;font-weight:600">Issue Type</td><td style="padding:8px 12px;color:#1e293b">{f.get('service_type','-')}</td></tr>
                  <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Description</td><td style="padding:8px 12px;color:#1e293b">{f.get('description','-')}</td></tr>
                </table>
                <p style="color:#6b7280;font-size:0.9rem">Log in to the Admin Portal to accept or decline this request.</p>
              </div>
              <div style="background:#f8fafc;padding:12px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">Qualesce IT Asset Tracker</div>
            </div>"""
            send_notification_email(admin_emails, f'New Service Request #{req_id} from {employee_name}', html)
        flash('Service request submitted! You will be notified once reviewed.', 'success')
        return redirect(url_for('user_my_requests'))
    assets = conn.execute(
        "SELECT asset_no, model FROM assets WHERE login_id LIKE ? ORDER BY asset_no",
        (f'%{employee_name}%',)).fetchall()
    conn.close()
    return render_template('user/request_service.html', assets=assets)


# ---------------------------------------------------------------------------
# Stock Dashboard import / template
# ---------------------------------------------------------------------------

@app.route('/admin/stock-import', methods=['POST'])
@admin_required
def admin_stock_import():
    file = request.files.get('excel_file')
    if not file or not file.filename.lower().endswith(('.xlsx', '.xls')):
        flash('Please upload a valid Excel file (.xlsx or .xls).', 'danger')
        return redirect(url_for('admin_dashboard'))
    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        sheet_name = 'Stock Dashboard' if 'Stock Dashboard' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]

        # Locate the "Row Labels" header row
        header_row_num = None
        for row in ws.iter_rows():
            for cell in row:
                if str(cell.value).strip() == 'Row Labels':
                    header_row_num = cell.row
                    break
            if header_row_num:
                break

        if not header_row_num:
            flash('Could not find "Row Labels" in the file. Please use the provided template.', 'danger')
            return redirect(url_for('admin_dashboard'))

        # Collect model names from header row (skip col 1 and "Grand Total")
        header_cells = list(ws.iter_rows(min_row=header_row_num, max_row=header_row_num, values_only=False))[0]
        models = [(cell.column, str(cell.value).strip())
                  for cell in header_cells[1:]
                  if cell.value and str(cell.value).strip() not in ('Grand Total', '')]

        # Parse data rows
        conn = get_db()
        conn.execute("DELETE FROM stock_dashboard")
        for row in ws.iter_rows(min_row=header_row_num + 1, values_only=False):
            status = str(row[0].value).strip() if row[0].value else ''
            if not status or status == 'Grand Total':
                continue
            for col_idx, model in models:
                cell = ws.cell(row=row[0].row, column=col_idx)
                try:
                    count = int(cell.value or 0)
                except (TypeError, ValueError):
                    count = 0
                if count > 0:
                    conn.execute(
                        "INSERT OR REPLACE INTO stock_dashboard (status, model, count) VALUES (?,?,?)",
                        (status, model, count))
        conn.commit()
        conn.close()
        flash(f'Stock Dashboard imported from "{sheet_name}" sheet successfully!', 'success')
    except Exception as e:
        flash(f'Import error: {e}', 'danger')
    return redirect(url_for('admin_dashboard'))


@app.route('/admin/stock-template')
@admin_required
def admin_stock_template():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Stock Dashboard'

    models = [
        'DELL Inspiron N5010', 'DELL LATITUDE 3400', 'DELL LATITUDE 3410',
        'Dell Latitude 5310', 'Dell LATITUDE 7480', 'DELL LATITUDE E7450',
        'DELL Vostro 3558', 'LENOVO THINKPAD T450', 'Lenovo B40-80',
        'LENOVO V130', 'LENOVO V310'
    ]
    statuses = ['Assigned', 'DEAD', 'IT Stock', 'Service', 'To Check']

    header_fill  = PatternFill('solid', fgColor='1F3864')
    total_fill   = PatternFill('solid', fgColor='D6E4F0')
    thin_border  = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'))

    # Row 3 — pivot label row
    ws['A3'] = 'Count of Resolution'
    ws['B3'] = 'Column Labels'

    # Row 4 — header
    ws.cell(row=4, column=1, value='Row Labels').font = Font(bold=True, color='FFFFFF')
    ws.cell(row=4, column=1).fill = header_fill
    ws.cell(row=4, column=1).alignment = Alignment(horizontal='center')
    for ci, model in enumerate(models, start=2):
        c = ws.cell(row=4, column=ci, value=model)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = header_fill
        c.alignment = Alignment(horizontal='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 16
    gt_col = len(models) + 2
    ws.cell(row=4, column=gt_col, value='Grand Total').font = Font(bold=True, color='FFFFFF')
    ws.cell(row=4, column=gt_col).fill = header_fill
    ws.cell(row=4, column=gt_col).alignment = Alignment(horizontal='center')
    ws.column_dimensions['A'].width = 14

    # Status rows
    for ri, status in enumerate(statuses, start=5):
        ws.cell(row=ri, column=1, value=status).font = Font(bold=True)
        ws.cell(row=ri, column=1).border = thin_border
        row_total = 0
        for ci, _ in enumerate(models, start=2):
            c = ws.cell(row=ri, column=ci, value=0)
            c.alignment = Alignment(horizontal='center')
            c.border = thin_border
        ws.cell(row=ri, column=gt_col, value=0).font = Font(bold=True)
        ws.cell(row=ri, column=gt_col).border = thin_border

    # Grand Total row
    gt_row = 5 + len(statuses)
    ws.cell(row=gt_row, column=1, value='Grand Total').font = Font(bold=True)
    ws.cell(row=gt_row, column=1).fill = total_fill
    for ci in range(2, gt_col + 1):
        ws.cell(row=gt_row, column=ci, value=0).fill = total_fill
        ws.cell(row=gt_row, column=ci).font = Font(bold=True)
        ws.cell(row=gt_row, column=ci).border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='stock_dashboard_template.xlsx'
    )


# ---------------------------------------------------------------------------
# Email configuration
# ---------------------------------------------------------------------------

@app.route('/admin/email-config', methods=['GET', 'POST'])
@admin_required
def admin_email_config():
    cfg = configparser.ConfigParser()
    cfg.read(CONFIG_PATH)

    if request.method == 'POST':
        f = request.form
        action = f.get('action', 'smtp')

        if action == 'smtp':
            if 'SMTP' not in cfg:
                cfg['SMTP'] = {}
            cfg['SMTP']['SMTP_HOST']     = f.get('smtp_host', 'smtp.gmail.com').strip()
            cfg['SMTP']['SMTP_PORT']     = f.get('smtp_port', '587').strip()
            cfg['SMTP']['SMTP_USER']     = f.get('smtp_user', '').strip()
            cfg['SMTP']['SMTP_PASSWORD'] = f.get('smtp_password', '').strip()
            cfg['SMTP']['FROM_NAME']     = f.get('from_name', 'Qualesce IT Tracker').strip()
            with open(CONFIG_PATH, 'w') as fout:
                cfg.write(fout)
            flash('SMTP settings saved successfully!', 'success')

        elif action == 'notifications':
            if 'NOTIFICATIONS' not in cfg:
                cfg['NOTIFICATIONS'] = {}
            notif_keys = [
                'notify_new_request', 'notify_employee_action',
                'notify_technician_assigned', 'notify_status_employee',
                'notify_status_admin', 'notify_chat',
            ]
            for key in notif_keys:
                cfg['NOTIFICATIONS'][key] = 'true' if f.get(key) == 'on' else 'false'
            with open(CONFIG_PATH, 'w') as fout:
                cfg.write(fout)
            flash('Notification preferences saved!', 'success')

        return redirect(url_for('admin_email_config'))

    smtp  = dict(cfg['SMTP'])          if 'SMTP'          in cfg else {}
    notif = dict(cfg['NOTIFICATIONS']) if 'NOTIFICATIONS' in cfg else {}
    return render_template('admin/email_config.html', smtp=smtp, notif=notif)


@app.route('/admin/test-email', methods=['POST'])
@admin_required
def admin_test_email():
    to_email = request.form.get('test_email', '').strip()
    if not to_email:
        flash('Please enter a test email address.', 'danger')
        return redirect(url_for('admin_email_config'))
    html = f"""
    <div style="font-family:Segoe UI,sans-serif;max-width:480px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
      <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:24px 30px;text-align:center">
        <h2 style="color:#fff;margin:0;font-size:1.2rem">Test Email — IT Asset Tracker</h2>
      </div>
      <div style="padding:28px 30px">
        <p style="color:#374151">Your email configuration is working correctly!</p>
        <p style="color:#6b7280;font-size:0.9rem">All system notifications (service requests, status updates, chat messages) will be delivered using this email account.</p>
        <div style="background:#f0fdf4;border:1px solid #bbf7d0;border-radius:8px;padding:14px 18px;margin-top:16px;color:#166534;font-size:0.9rem">
          <strong>✓ SMTP connection successful</strong><br>
          Sent at {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
        </div>
      </div>
      <div style="background:#f8fafc;padding:12px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">Qualesce IT Asset Tracker</div>
    </div>"""
    ok, msg = send_notification_email([to_email], 'IT Asset Tracker — Test Email', html)
    if ok:
        flash(f'Test email sent to {to_email} successfully!', 'success')
    else:
        flash(f'Failed to send test email: {msg}', 'danger')
    return redirect(url_for('admin_email_config'))


# ---------------------------------------------------------------------------
# API endpoints
# ---------------------------------------------------------------------------

@app.route('/api/assets/search')
@admin_required
def api_search_assets():
    q = request.args.get('q', '')
    conn = get_db()
    assets = conn.execute(
        "SELECT asset_no, login_id, model, asset_status FROM assets WHERE asset_no LIKE ? OR login_id LIKE ? LIMIT 20",
        (f'%{q}%', f'%{q}%')).fetchall()
    conn.close()
    return jsonify([dict(a) for a in assets])


# ---------------------------------------------------------------------------
# Personal email preferences - User & Technician
# ---------------------------------------------------------------------------

@app.route('/user/email-preferences', methods=['GET', 'POST'])
@login_required
def user_email_preferences():
    if session.get('role') in ('admin', 'technician'):
        return redirect(url_for('admin_dashboard'))
    user_id = session['user_id']
    conn = get_db()
    if request.method == 'POST':
        prefs = {
            'action': request.form.get('notify_action') == 'on',
            'status': request.form.get('notify_status') == 'on',
            'chat':   request.form.get('notify_chat')   == 'on',
        }
        conn.execute("UPDATE users SET notification_prefs=? WHERE id=?",
                     (json.dumps(prefs), user_id))
        conn.commit()
        conn.close()
        flash('Email preferences saved!', 'success')
        return redirect(url_for('user_email_preferences'))
    row = conn.execute("SELECT notification_prefs, email FROM users WHERE id=?", (user_id,)).fetchone()
    conn.close()
    try:
        prefs = json.loads(row['notification_prefs']) if row and row['notification_prefs'] else {}
    except Exception:
        prefs = {}
    user_email = row['email'] if row else ''
    return render_template('user/email_preferences.html', prefs=prefs, user_email=user_email)


@app.route('/technician/email-preferences', methods=['GET', 'POST'])
@technician_required
def technician_email_preferences():
    user_id = session['user_id']
    conn = get_db()
    if request.method == 'POST':
        prefs = {
            'assigned': request.form.get('notify_assigned') == 'on',
            'chat':     request.form.get('notify_chat')     == 'on',
            'status':   request.form.get('notify_status')   == 'on',
        }
        conn.execute("UPDATE users SET notification_prefs=? WHERE id=?",
                     (json.dumps(prefs), user_id))
        conn.commit()
        conn.close()
        flash('Email preferences saved!', 'success')
        return redirect(url_for('technician_email_preferences'))
    row = conn.execute("SELECT notification_prefs, email FROM users WHERE id=?", (user_id,)).fetchone()
    conn.close()
    try:
        prefs = json.loads(row['notification_prefs']) if row and row['notification_prefs'] else {}
    except Exception:
        prefs = {}
    user_email = row['email'] if row else ''
    return render_template('technician/email_preferences.html', prefs=prefs, user_email=user_email)


# ---------------------------------------------------------------------------
# Service Requests - User routes
# ---------------------------------------------------------------------------

@app.route('/user/my-requests')
@login_required
def user_my_requests():
    if session.get('role') in ('admin', 'technician'):
        return redirect(url_for('admin_dashboard'))
    conn = get_db()
    reqs = conn.execute(
        "SELECT * FROM service_requests WHERE employee_id=? ORDER BY created_at DESC",
        (session['user_id'],)).fetchall()
    conn.close()
    return render_template('user/my_requests.html', requests=reqs)


@app.route('/user/request/<int:req_id>', methods=['GET', 'POST'])
@login_required
def user_request_detail(req_id):
    if session.get('role') in ('admin', 'technician'):
        return redirect(url_for('admin_dashboard'))
    conn = get_db()
    req = conn.execute(
        "SELECT * FROM service_requests WHERE id=? AND employee_id=?",
        (req_id, session['user_id'])).fetchone()
    if not req:
        conn.close()
        flash('Request not found.', 'danger')
        return redirect(url_for('user_my_requests'))
    if request.method == 'POST':
        message = request.form.get('message', '').strip()
        if message:
            conn.execute(
                "INSERT INTO chat_messages (request_id, sender_id, sender_name, sender_role, message) VALUES (?,?,?,?,?)",
                (req_id, session['user_id'], session['employee_name'], 'user', message))
            conn.commit()
            admins = conn.execute("SELECT email FROM users WHERE role='admin' AND email!=''").fetchall()
            admin_emails = [a['email'] for a in admins if a['email']]
            if admin_emails and get_notification_prefs()['chat']:
                send_notification_email(admin_emails, f'New message on Request #{req_id}',
                    build_chat_email_html(session['employee_name'], 'user', message, req_id, req['service_type']))
            if req['assigned_to_id']:
                tech = conn.execute("SELECT email FROM users WHERE id=?", (req['assigned_to_id'],)).fetchone()
                if _should_notify_tech(req['assigned_to_id'], tech['email'] if tech else '', 'chat', 'chat'):
                    send_notification_email([tech['email']], f'New message on Request #{req_id}',
                        build_chat_email_html(session['employee_name'], 'user', message, req_id, req['service_type']))
        conn.close()
        return redirect(url_for('user_request_detail', req_id=req_id))
    messages = conn.execute(
        "SELECT * FROM chat_messages WHERE request_id=? ORDER BY created_at ASC",
        (req_id,)).fetchall()
    conn.close()
    return render_template('user/request_detail.html', req=req, messages=messages)


# ---------------------------------------------------------------------------
# Service Requests - Admin routes
# ---------------------------------------------------------------------------

@app.route('/admin/service-requests')
@admin_required
def admin_service_requests():
    status_f = request.args.get('status', '')
    conn = get_db()
    q = "SELECT * FROM service_requests WHERE 1=1"
    p = []
    if status_f:
        q += " AND status=?"
        p.append(status_f)
    q += " ORDER BY created_at DESC"
    reqs = conn.execute(q, p).fetchall()
    pending_count = conn.execute("SELECT COUNT(*) FROM service_requests WHERE status='Pending'").fetchone()[0]
    conn.close()
    return render_template('admin/service_requests.html', requests=reqs, status_filter=status_f, pending_count=pending_count)


@app.route('/admin/service-request/<int:req_id>')
@admin_required
def admin_service_request_detail(req_id):
    conn = get_db()
    req = conn.execute("SELECT * FROM service_requests WHERE id=?", (req_id,)).fetchone()
    if not req:
        conn.close()
        flash('Request not found.', 'danger')
        return redirect(url_for('admin_service_requests'))
    messages = conn.execute(
        "SELECT * FROM chat_messages WHERE request_id=? ORDER BY created_at ASC",
        (req_id,)).fetchall()
    technicians = conn.execute(
        "SELECT id, employee_name, username FROM users WHERE role='technician' ORDER BY employee_name").fetchall()
    conn.close()
    return render_template('admin/service_request_detail.html', req=req, messages=messages, technicians=technicians)


@app.route('/admin/service-request/<int:req_id>/action', methods=['POST'])
@admin_required
def admin_service_request_action(req_id):
    action = request.form.get('action')
    new_status = 'Accepted' if action == 'accept' else 'Declined'
    conn = get_db()
    req = conn.execute("SELECT * FROM service_requests WHERE id=?", (req_id,)).fetchone()
    if not req:
        conn.close()
        flash('Request not found.', 'danger')
        return redirect(url_for('admin_service_requests'))
    conn.execute("UPDATE service_requests SET status=?, updated_at=? WHERE id=?",
                 (new_status, datetime.now().isoformat(), req_id))
    conn.commit()
    emp = conn.execute("SELECT email FROM users WHERE id=?", (req['employee_id'],)).fetchone()
    conn.close()
    if _should_notify_employee(req['employee_id'], emp['email'] if emp else '', 'employee_action', 'action'):
        color = '#22c55e' if action == 'accept' else '#ef4444'
        html = f"""
        <div style="font-family:Segoe UI,sans-serif;max-width:560px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
          <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:24px 30px;text-align:center">
            <h2 style="color:#fff;margin:0;font-size:1.2rem">Service Request #{req_id} — {new_status}</h2>
          </div>
          <div style="padding:28px 30px">
            <p style="color:#374151">Hi <strong>{req['employee_name']}</strong>,</p>
            <p style="color:#374151">Your service request has been <strong style="color:{color}">{new_status}</strong>.</p>
            <table style="width:100%;border-collapse:collapse;margin:16px 0">
              <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;width:40%;font-weight:600">Request #</td><td style="padding:8px 12px;color:#1e293b">#{req_id}</td></tr>
              <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Asset</td><td style="padding:8px 12px;color:#1e293b">{req['asset_no']}</td></tr>
              <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;font-weight:600">Issue Type</td><td style="padding:8px 12px;color:#1e293b">{req['service_type']}</td></tr>
              <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Status</td><td style="padding:8px 12px;font-weight:700;color:{color}">{new_status}</td></tr>
            </table>
          </div>
          <div style="background:#f8fafc;padding:12px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">Qualesce IT Asset Tracker</div>
        </div>"""
        send_notification_email([emp['email']], f'Your Service Request #{req_id} has been {new_status}', html)
    flash(f'Request {new_status.lower()}!', 'success')
    return redirect(url_for('admin_service_request_detail', req_id=req_id))


@app.route('/admin/service-request/<int:req_id>/status', methods=['POST'])
@admin_required
def admin_service_request_status(req_id):
    new_status = request.form.get('status')
    if new_status not in ('Pending', 'Accepted', 'In Progress', 'Hold', 'Completed', 'Declined'):
        flash('Invalid status.', 'danger')
        return redirect(url_for('admin_service_request_detail', req_id=req_id))
    conn = get_db()
    req = conn.execute("SELECT * FROM service_requests WHERE id=?", (req_id,)).fetchone()
    conn.execute("UPDATE service_requests SET status=?, updated_at=? WHERE id=?",
                 (new_status, datetime.now().isoformat(), req_id))
    conn.commit()
    emp = conn.execute("SELECT email FROM users WHERE id=?", (req['employee_id'],)).fetchone()
    conn.close()
    if _should_notify_employee(req['employee_id'], emp['email'] if emp else '', 'status_employee', 'status'):
        send_notification_email([emp['email']], f'Service Request #{req_id} Status: {new_status}',
            build_status_email_html(req['employee_name'], new_status, req_id, req['service_type'], req['asset_no']))
    flash(f'Status updated to {new_status}!', 'success')
    return redirect(url_for('admin_service_request_detail', req_id=req_id))


@app.route('/admin/service-request/<int:req_id>/assign', methods=['POST'])
@admin_required
def admin_service_request_assign(req_id):
    tech_id = request.form.get('technician_id')
    conn = get_db()
    tech = conn.execute(
        "SELECT id, employee_name, username, email FROM users WHERE id=? AND role='technician'",
        (tech_id,)).fetchone()
    if not tech:
        conn.close()
        flash('Technician not found.', 'danger')
        return redirect(url_for('admin_service_request_detail', req_id=req_id))
    req = conn.execute("SELECT * FROM service_requests WHERE id=?", (req_id,)).fetchone()
    tech_name = tech['employee_name'] or tech['username']
    conn.execute(
        "UPDATE service_requests SET assigned_to_id=?, assigned_to_name=?, status='In Progress', updated_at=? WHERE id=?",
        (tech['id'], tech_name, datetime.now().isoformat(), req_id))
    conn.commit()
    prefs = get_notification_prefs()
    if _should_notify_tech(tech['id'], tech['email'], 'technician_assigned', 'assigned'):
        html = f"""
        <div style="font-family:Segoe UI,sans-serif;max-width:560px;margin:auto;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden">
          <div style="background:linear-gradient(135deg,#0f3460,#533483);padding:24px 30px;text-align:center">
            <h2 style="color:#fff;margin:0;font-size:1.2rem">New Task Assigned — Request #{req_id}</h2>
          </div>
          <div style="padding:28px 30px">
            <p style="color:#374151">Hi <strong>{tech_name}</strong>, a service request has been assigned to you.</p>
            <table style="width:100%;border-collapse:collapse;margin:16px 0">
              <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;width:40%;font-weight:600">Employee</td><td style="padding:8px 12px;color:#1e293b">{req['employee_name']}</td></tr>
              <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Asset</td><td style="padding:8px 12px;color:#1e293b">{req['asset_no']}</td></tr>
              <tr><td style="padding:8px 12px;background:#f8fafc;color:#6b7280;font-weight:600">Issue</td><td style="padding:8px 12px;color:#1e293b">{req['service_type']}</td></tr>
              <tr><td style="padding:8px 12px;background:#f1f5f9;color:#6b7280;font-weight:600">Description</td><td style="padding:8px 12px;color:#1e293b">{req['description']}</td></tr>
            </table>
            <p style="color:#6b7280;font-size:0.9rem">Log in to your Technician Portal to view and respond.</p>
          </div>
          <div style="background:#f8fafc;padding:12px 30px;text-align:center;color:#94a3b8;font-size:0.8rem">Qualesce IT Asset Tracker</div>
        </div>"""
        send_notification_email([tech['email']], f'Service Request #{req_id} Assigned to You', html)
    emp = conn.execute("SELECT email FROM users WHERE id=?", (req['employee_id'],)).fetchone()
    conn.close()
    if _should_notify_employee(req['employee_id'], emp['email'] if emp else '', 'status_employee', 'status'):
        send_notification_email([emp['email']], f'Your Request #{req_id} is Now In Progress',
            build_status_email_html(req['employee_name'], 'In Progress', req_id, req['service_type'], req['asset_no']))
    flash(f'Request assigned to {tech_name} and set to In Progress!', 'success')
    return redirect(url_for('admin_service_request_detail', req_id=req_id))


@app.route('/admin/service-request/<int:req_id>/chat', methods=['POST'])
@admin_required
def admin_service_request_chat(req_id):
    message = request.form.get('message', '').strip()
    if not message:
        return redirect(url_for('admin_service_request_detail', req_id=req_id))
    conn = get_db()
    req = conn.execute("SELECT * FROM service_requests WHERE id=?", (req_id,)).fetchone()
    conn.execute(
        "INSERT INTO chat_messages (request_id, sender_id, sender_name, sender_role, message) VALUES (?,?,?,?,?)",
        (req_id, session['user_id'], session['employee_name'], 'admin', message))
    conn.commit()
    emp = conn.execute("SELECT email FROM users WHERE id=?", (req['employee_id'],)).fetchone()
    if _should_notify_employee(req['employee_id'], emp['email'] if emp else '', 'chat', 'chat'):
        send_notification_email([emp['email']], f'Admin replied on Request #{req_id}',
            build_chat_email_html(session['employee_name'], 'admin', message, req_id, req['service_type']))
    if req['assigned_to_id']:
        tech = conn.execute("SELECT email FROM users WHERE id=?", (req['assigned_to_id'],)).fetchone()
        if _should_notify_tech(req['assigned_to_id'], tech['email'] if tech else '', 'chat', 'chat'):
            send_notification_email([tech['email']], f'New message on Request #{req_id}',
                build_chat_email_html(session['employee_name'], 'admin', message, req_id, req['service_type']))
    conn.close()
    return redirect(url_for('admin_service_request_detail', req_id=req_id))


# ---------------------------------------------------------------------------
# Technician routes
# ---------------------------------------------------------------------------

@app.route('/technician')
@technician_required
def technician_dashboard():
    conn = get_db()
    reqs = conn.execute(
        "SELECT * FROM service_requests WHERE assigned_to_id=? ORDER BY updated_at DESC",
        (session['user_id'],)).fetchall()
    stats = {
        'total': len(reqs),
        'in_progress': sum(1 for r in reqs if r['status'] == 'In Progress'),
        'hold': sum(1 for r in reqs if r['status'] == 'Hold'),
        'completed': sum(1 for r in reqs if r['status'] == 'Completed'),
    }
    conn.close()
    return render_template('technician/dashboard.html', requests=reqs, stats=stats)


@app.route('/technician/request/<int:req_id>', methods=['GET', 'POST'])
@technician_required
def technician_request_detail(req_id):
    conn = get_db()
    req = conn.execute(
        "SELECT * FROM service_requests WHERE id=? AND assigned_to_id=?",
        (req_id, session['user_id'])).fetchone()
    if not req:
        conn.close()
        flash('Request not found.', 'danger')
        return redirect(url_for('technician_dashboard'))
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'chat':
            message = request.form.get('message', '').strip()
            if message:
                conn.execute(
                    "INSERT INTO chat_messages (request_id, sender_id, sender_name, sender_role, message) VALUES (?,?,?,?,?)",
                    (req_id, session['user_id'], session['employee_name'], 'technician', message))
                conn.commit()
                emp = conn.execute("SELECT email FROM users WHERE id=?", (req['employee_id'],)).fetchone()
                admins = conn.execute("SELECT email FROM users WHERE role='admin' AND email!=''").fetchall()
                recipients = []
                if emp and emp['email']:
                    recipients.append(emp['email'])
                recipients += [a['email'] for a in admins if a['email']]
                if get_notification_prefs()['chat']:
                    if emp and emp['email'] and get_user_notif_prefs(req['employee_id']).get('chat', True):
                        send_notification_email([emp['email']], f'Technician message on Request #{req_id}',
                            build_chat_email_html(session['employee_name'], 'technician', message, req_id, req['service_type']))
                    admin_emails_chat = [a['email'] for a in admins if a['email']]
                    if admin_emails_chat:
                        send_notification_email(admin_emails_chat, f'Technician message on Request #{req_id}',
                            build_chat_email_html(session['employee_name'], 'technician', message, req_id, req['service_type']))
        elif action == 'status':
            new_status = request.form.get('status')
            if new_status in ('In Progress', 'Hold', 'Completed'):
                conn.execute("UPDATE service_requests SET status=?, updated_at=? WHERE id=?",
                             (new_status, datetime.now().isoformat(), req_id))
                conn.commit()
                emp = conn.execute("SELECT email FROM users WHERE id=?", (req['employee_id'],)).fetchone()
                admins = conn.execute("SELECT email FROM users WHERE role='admin' AND email!=''").fetchall()
                _prefs = get_notification_prefs()
                if _should_notify_employee(req['employee_id'], emp['email'] if emp else '', 'status_employee', 'status'):
                    send_notification_email([emp['email']], f'Request #{req_id} Status: {new_status}',
                        build_status_email_html(req['employee_name'], new_status, req_id, req['service_type'], req['asset_no']))
                admin_emails = [a['email'] for a in admins if a['email']]
                if admin_emails and _prefs['status_admin']:
                    send_notification_email(admin_emails, f'Technician updated Request #{req_id} to {new_status}',
                        build_status_email_html(req['employee_name'], new_status, req_id, req['service_type'], req['asset_no']))
                flash(f'Status updated to {new_status}!', 'success')
        conn.close()
        return redirect(url_for('technician_request_detail', req_id=req_id))
    messages = conn.execute(
        "SELECT * FROM chat_messages WHERE request_id=? ORDER BY created_at ASC",
        (req_id,)).fetchall()
    conn.close()
    return render_template('technician/request_detail.html', req=req, messages=messages)


if __name__ == '__main__':
    init_db()
    import_excel()
    print("\n" + "="*50)
    print("  IT Asset Tracker is running!")
    print("  Open: http://localhost:5000")
    print("  Admin: admin / admin123")
    print("  User:  (employee username) / pass123")
    print("="*50 + "\n")
    app.run(debug=False, host='0.0.0.0', port=5000, use_reloader=False)
